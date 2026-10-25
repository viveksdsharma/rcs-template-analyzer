"""PowerPoint builder (python-pptx). On the Vi template when one is supplied; on a plain design otherwise.
Flow follows the Axis deck: template in focus -> what might not work -> recommendations, plus period comparison."""
import os

from lxml import etree
from PIL import Image
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

from . import metrics as M
from .render_docx import YOY_DEC, YOY_ROWS, _change_text

RED, NAVY, PURPLE, TXT, MUTED = "EE2737", "2F3043", "5F004B", "2A2A33", "6B6B76"
PANEL, YEL, WHITE = "F2F2F5", "FFC600", "FFFFFF"


def rgb(h):
    return RGBColor.from_string(h)


class Deck:
    """Wraps either the Vi template (sample slides removed) or a plain 16:9 presentation."""

    def __init__(self, template_path=None):
        self.templated = bool(template_path and os.path.exists(template_path))
        if self.templated:
            self.prs = Presentation(template_path)
            lst = self.prs.slides._sldIdLst
            for sid in list(lst):
                self.prs.part.drop_rel(sid.rId)
                lst.remove(sid)
            self.LAY = {l.name.strip(): l for l in self.prs.slide_layouts}
        else:
            self.prs = Presentation()
            self.prs.slide_width, self.prs.slide_height = Inches(13.333), Inches(7.5)
            self.LAY = {}

    def slide(self):
        if self.templated:
            s = self.prs.slides.add_slide(self.LAY["Only Title slide"])
            for ph in list(s.placeholders):
                ph._element.getparent().remove(ph._element)
            return s
        return self.prs.slides.add_slide(self.prs.slide_layouts[6])

    def title_slide(self, title, brand, subtitle):
        if self.templated:
            s = self.prs.slides.add_slide(self.LAY["3_Title Slide 01"])
            for ph in s.placeholders:
                ph.text_frame.text = title if ph.placeholder_format.idx == 0 else str(brand).upper()
            text(s, 0.56, 5.35, 4.8, 0.4, subtitle, size=14, color=WHITE)
            return s
        s = self.prs.slides.add_slide(self.prs.slide_layouts[6])
        rect(s, 0, 0, 13.333, 7.5, PURPLE)
        rect(s, 8.6, 0, 4.733, 7.5, RED)
        text(s, 0.7, 2.0, 7.6, 1.6, title, size=40, bold=True, color=WHITE)
        text(s, 0.7, 3.9, 7.6, 0.5, str(brand).upper(), size=22, color="F3DCE4")
        text(s, 0.7, 4.6, 7.6, 0.4, subtitle, size=14, color=WHITE)
        return s

    def thanks_slide(self):
        if self.templated:
            s = self.prs.slides.add_slide(self.LAY["1_Thank You"])
            for ph in list(s.placeholders):
                if ph.placeholder_format.idx == 0:
                    ph._element.getparent().remove(ph._element)
            return s
        s = self.prs.slides.add_slide(self.prs.slide_layouts[6])
        rect(s, 0, 0, 13.333, 7.5, PURPLE)
        text(s, 0.7, 3.0, 11, 1.2, "thank you", size=44, bold=True, color=WHITE)
        return s


def rect(s, x, y, w, h, fill, line=None, shape=MSO_SHAPE.RECTANGLE):
    sh = s.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.fill.solid()
    sh.fill.fore_color.rgb = rgb(fill)
    if line:
        sh.line.color.rgb = rgb(line)
        sh.line.width = Pt(0.75)
    else:
        sh.line.fill.background()
    sh.shadow.inherit = False
    return sh


def _bullet(p, ind=0.18):
    pPr = p._p.get_or_add_pPr()
    pPr.set("marL", str(int(Inches(ind))))
    pPr.set("indent", str(-int(Inches(ind))))
    for t in ("a:buNone", "a:buChar", "a:buAutoNum"):
        for e in pPr.findall(qn(t)):
            pPr.remove(e)
    etree.SubElement(pPr, qn("a:buChar")).set("char", "\u2022")


def text(s, x, y, w, h, content, size=12, bold=False, color=TXT, align="l", anchor="t", bullets=False, space_after=0, italic=False):
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE}[anchor]
    paras = content if isinstance(content, list) else [content]
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT}[align]
        if space_after:
            p.space_after = Pt(space_after)
        para_bullet = False
        if isinstance(para, dict):
            para_bullet = bool(para.get("bullet"))
            para = para["runs"]
        if isinstance(para, list):
            runs = para
        else:
            runs = [(x, {"bold": True} if k % 2 else {}) for k, x in enumerate(str(para).split("**")) if x]
        for t, o in runs:
            r = p.add_run()
            r.text = t
            r.font.size = Pt(o.get("size", size))
            r.font.bold = o.get("bold", bold)
            r.font.italic = o.get("italic", italic)
            r.font.color.rgb = rgb(o.get("color", color))
        if bullets or para_bullet:
            _bullet(p)
    return tb


def panel_text(items, size=12.5):
    out = []
    for it in items:
        out.append([(it["h"], {"bold": True, "color": RED, "size": size + 1.5})])
        out.append(it["b"])
    return out


def header(s, eyebrow, title, sub=None):
    text(s, 0.5, 0.32, 10.2, 0.3, eyebrow, size=16, bold=True, color=RED)
    text(s, 0.5, 0.58, 10.3, 0.55, title, size=26, bold=True, color=NAVY)
    if sub:
        text(s, 0.5, 1.05, 10.3, 0.3, sub, size=13, color=MUTED, italic=True)
    rect(s, 0.5, 1.40 if sub else 1.18, 12.3, 0.03, RED)


def source(s, t):
    text(s, 3.6, 6.88, 9.2, 0.25, t, size=9, color=MUTED, align="r")


def notes(s, t):
    if t:
        s.notes_slide.notes_text_frame.text = t


def table(s, x, y, w, colw, rows, size=11, rowh=0.36, align=None):
    nr, nc = len(rows), len(rows[0])
    gs = s.shapes.add_table(nr, nc, Inches(x), Inches(y), Inches(w), Inches(rowh * nr))
    tbl = gs.table
    tbl._tbl.tblPr.set("firstRow", "0")
    tbl._tbl.tblPr.set("bandRow", "0")
    for i, cw in enumerate(colw):
        tbl.columns[i].width = Inches(cw)
    for r in range(nr):
        tbl.rows[r].height = Inches(rowh)
        for c in range(nc):
            cell = tbl.cell(r, c)
            v = rows[r][c]
            opts = {}
            if isinstance(v, tuple):
                v, opts = v
            cell.margin_left = cell.margin_right = Inches(0.08)
            cell.margin_top = cell.margin_bottom = Inches(0.03)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.fill.solid()
            cell.fill.fore_color.rgb = rgb(opts.get("fill") or (NAVY if r == 0 else ("FFFFFF" if r % 2 else PANEL)))
            p = cell.text_frame.paragraphs[0]
            cell.text_frame.word_wrap = True
            p.alignment = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT}[(align[c] if align else "l")]
            run = p.add_run()
            run.text = str(v)
            run.font.size = Pt(opts.get("size", size))
            run.font.bold = opts.get("bold", r == 0)
            run.font.color.rgb = rgb(opts.get("color", "FFFFFF" if r == 0 else TXT))
    return tbl


def bar_chart(s, x, y, w, h, cats, series, colors=None, point_colors=None, horizontal=False, maxv=None, fmt='0.00"%"', legend=False, font=11, gap=60, label_size=11):
    cd = CategoryChartData()
    cd.categories = cats
    for name, vals in series:
        cd.add_series(name, vals)
    gf = s.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED if horizontal else XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(x), Inches(y), Inches(w), Inches(h), cd)
    ch = gf.chart
    ch.font.size = Pt(font)
    ch.font.color.rgb = rgb(TXT)
    ch.has_title = False
    ch.has_legend = legend
    if legend:
        ch.legend.position = XL_LEGEND_POSITION.BOTTOM
        ch.legend.include_in_layout = False
        ch.legend.font.size = Pt(font)
    va, ca = ch.value_axis, ch.category_axis
    va.visible = False
    va.has_major_gridlines = False
    va.minimum_scale = 0
    if maxv:
        va.maximum_scale = maxv
    ca.format.line.color.rgb = rgb("BFBFC6")
    ca.tick_labels.font.size = Pt(font)
    if horizontal:
        ca.reverse_order = True
    plot = ch.plots[0]
    plot.gap_width = gap
    plot.has_data_labels = True
    dl = plot.data_labels
    dl.number_format = fmt
    dl.number_format_is_linked = False
    dl.position = XL_LABEL_POSITION.OUTSIDE_END
    dl.font.size = Pt(label_size)
    dl.font.bold = True
    dl.font.color.rgb = rgb(NAVY)
    cols = colors or [NAVY, RED]
    for i, ser in enumerate(plot.series):
        ser.format.fill.solid()
        ser.format.fill.fore_color.rgb = rgb(cols[i % len(cols)])
        if point_colors and i == 0:
            for j, pc in enumerate(point_colors):
                pt = ser.points[j]
                pt.format.fill.solid()
                pt.format.fill.fore_color.rgb = rgb(pc)
    return ch


def picture(s, path, x, y, max_w, max_h):
    with Image.open(path) as im_:
        iw, ih = im_.size
    ratio = iw / ih
    w = max_w
    h = w / ratio
    if h > max_h:
        h = max_h
        w = h * ratio
    s.shapes.add_picture(path, Inches(x), Inches(y), Inches(w), Inches(h))
    box = rect(s, x, y, w, h, WHITE)
    box.fill.background()
    box.line.color.rgb = rgb("D9D9E0")
    box.line.width = Pt(0.75)
    return w, h


def build_pptx(metrics, analysis, request, template_path, out_path, style, image_dir):
    D = Deck(template_path)
    brand = request.get("brand_display") or request.get("brand")
    labels = request.get("period_labels", {})
    ylab = lambda y: labels.get(str(y), str(y))  # noqa: E731
    years = metrics["meta"]["years"]
    cur = max(years)
    dk = analysis.get("deck", {}) or {}
    tpl = {t["template"]: t for t in metrics["templates"]}
    req_t = {t["template"]: t for t in request["templates"]}
    ana_t = {t["template"]: t for t in analysis["templates"]}
    SRC = "Source: partner portal template export. * = our calculation from the sheet counts."
    shorts = {x["template"]: x.get("short") for x in (request["templates"] + (request.get("reference_templates") or []))}
    sr = lambda t, k: M.shown_rate(t, k)  # noqa: E731

    D.title_slide(analysis["meta"].get("deck_title", "RCS Template Analysis"), brand, analysis["meta"].get("deck_subtitle") or request.get("period", ""))

    ctx = request.get("context")
    if ctx and ctx.get("stats"):
        s = D.slide()
        header(s, "CONTEXT", "Why we are here")
        text(s, 0.5, 1.45, 6.2, 0.3, ctx.get("stats_title", f"What {brand} reported"), size=14, bold=True, color=NAVY)
        for i, st in enumerate(ctx["stats"][:4]):
            cx, cy = 0.5 + (i % 2) * 3.15, 1.9 + (i // 2) * 2.15
            rect(s, cx, cy, 3.0, 1.95, PANEL)
            rect(s, cx, cy, 0.06, 1.95, RED)
            text(s, cx + 0.25, cy + 0.2, 2.6, 0.7, st["big"], size=26 if len(st["big"]) < 8 else 21, bold=True, color=NAVY, anchor="m")
            text(s, cx + 0.25, cy + 1.0, 2.6, 0.8, st["label"].split("\n"), size=12, color=MUTED, space_after=2)
        rect(s, 7.1, 1.45, 5.7, 5.0, PANEL)
        rect(s, 7.1, 1.45, 5.7, 0.5, NAVY)
        text(s, 7.35, 1.45, 5.3, 0.5, "What was agreed, and what this deck does", size=15, bold=True, color=WHITE, anchor="m")
        text(s, 7.35, 2.15, 5.2, 4.2, [[(a["label"] + "  ", {"bold": True, "color": RED}), (a["text"], {})] for a in ctx.get("agreed", [])[:5]], size=13, space_after=12)
        source(s, ctx.get("source", "Source: figures as reported by the brand."))

    s = D.slide()
    header(s, "WHAT WE ANALYSED", f"{len(metrics['templates'])} templates" + (f", {len(years)} periods" if len(years) > 1 else ""))
    rows = [["Year", "Template", "Date", "GoogleID", "Messages Submitted", "DeliveryRate", "ReadRate", "P2AResponseRate (click rate)"]]
    for t in metrics["templates"]:
        rows.append([str(t["year"]), t["template"], M.nice_date(t["date"]).rsplit(" ", 1)[0], t["google_id_short"], M.fmt_int(t["counts"]["submitted"]),
                     M.fmt_pct(sr(t, "DeliveryRate")), M.fmt_pct(sr(t, "ReadRate")), (M.fmt_pct(sr(t, "P2AResponseRate")), {"bold": True})])
    rh = min(0.42, 4.5 / len(rows))
    table(s, 0.5, 1.5, 12.3, [0.7, 3.2, 1.0, 1.2, 1.7, 1.4, 1.2, 1.9], rows, size=12 if len(rows) <= 9 else 10, rowh=rh, align=["l", "l", "l", "l", "r", "r", "r", "r"])
    text(s, 0.5, 1.5 + rh * len(rows) + 0.2, 12.3, 1.2, [
        [("Column names are the sheet's own.  ", {"bold": True, "color": NAVY}), ("DeliveryRate = delivered \u00F7 sent. ReadRate = read \u00F7 delivered. P2AResponseRate = click rate (responses \u00F7 read).", {})],
        [("Each template is a single send date.  ", {"bold": True, "color": NAVY}), ("Audience, day of week and frequency cannot be separated from the creative, so findings are directional.", {})]], size=12, space_after=5)
    source(s, SRC)

    s = D.slide()
    header(s, "THE FUNNEL", dk.get("funnel_title", "Where messages are lost along the way"))
    text(s, 0.5, 1.45, 7.2, 0.3, "Share of Messages Submitted", size=14, bold=True, color=NAVY)
    ser = []
    for y in sorted(years):
        g = metrics["groups"]["year"][str(y)]
        c = g["counts"]
        sub = c["submitted"]
        ser.append((f"{ylab(y)} ({g['n_templates']} templates)", [round(100 * c["failed"] / sub, 1), round(100 * c["sent"] / sub, 1), round(100 * c["delivered"] / sub, 1), round(100 * c["read"] / sub, 1)]))
    bar_chart(s, 0.4, 1.8, 7.4, 4.6, ["Failed pre-send", "Sent", "Delivered", "Read"], ser, colors=[PURPLE, RED], maxv=70, fmt='0.0"%"', legend=len(ser) > 1, font=12, gap=70)
    rect(s, 8.1, 1.45, 4.7, 5.0, PANEL)
    text(s, 8.35, 1.6, 4.3, 4.8, panel_text(dk.get("funnel_panel", [])), size=12.5, space_after=7)
    source(s, SRC)

    yoy = metrics["comparisons"].get("year_over_year")
    if yoy:
        s = D.slide()
        header(s, "PERIOD COMPARISON", dk.get("yoy_title", "What changed between the periods"))
        a, b = yoy["from_year"], yoy["to_year"]
        ga, gb = metrics["groups"]["year"][str(a)], metrics["groups"]["year"][str(b)]
        text(s, 0.5, 1.45, 5.6, 0.3, "Click rate (P2AResponseRate)", size=14, bold=True, color=NAVY)
        bar_chart(s, 0.4, 1.8, 5.5, 4.0, [f"{ylab(a)}\n({ga['n_templates']} templates)", f"{ylab(b)}\n({gb['n_templates']} templates)"],
                  [("Click rate", [round(ga["rates"]["P2AResponseRate"], 2), round(gb["rates"]["P2AResponseRate"], 2)])], point_colors=[PURPLE, RED],
                  maxv=max(ga["rates"]["P2AResponseRate"], gb["rates"]["P2AResponseRate"]) * 1.25, fmt='0.00"%"', font=12, gap=55, label_size=16)
        rows = [["Metric", ylab(a), ylab(b), "Change"]]
        for key, label in YOY_ROWS:
            m = yoy["metrics"].get(key)
            if m:
                hot = abs(m["change_pts"]) > 5 or (m["fold"] and (m["fold"] > 1.5 or m["fold"] < 0.67))
                d_ = YOY_DEC.get(key, 2)
                rows.append([label, M.fmt_pct(m["from"], d_), M.fmt_pct(m["to"], d_), (_change_text(key, m), {"bold": hot, "color": RED if hot else TXT})])
        table(s, 6.2, 1.5, 6.6, [2.5, 1.2, 1.2, 1.7], rows, size=12, rowh=0.47, align=["l", "r", "r", "r"])
        if dk.get("yoy_note"):
            yy = 1.5 + 0.47 * len(rows) + 0.2
            rect(s, 6.2, yy, 6.6, 1.5, PANEL)
            rect(s, 6.2, yy, 0.06, 1.5, RED)
            text(s, 6.45, yy + 0.1, 6.2, 1.3, dk["yoy_note"], size=12, anchor="m")
        source(s, SRC)

    s = D.slide()
    header(s, "EVERY TEMPLATE", dk.get("ranking_title", "Click rate by template"))
    ranked = sorted(metrics["templates"], key=lambda t: -t["rates"]["P2AResponseRate"])
    pal = {y: c for y, c in zip(sorted(years), [PURPLE, RED, NAVY])}
    text(s, 0.5, 1.45, 7.4, 0.3, "Click rate (P2AResponseRate)", size=14, bold=True, color=NAVY)
    bar_chart(s, 0.4, 1.8, 7.7, 4.45, [f"{shorts.get(t['template']) or t['template']} ({M.nice_date(t['date']) if len(years) > 1 else M.nice_date(t['date']).rsplit(' ', 1)[0]})" for t in ranked],
              [("Click rate", [round(sr(t, "P2AResponseRate"), 2) for t in ranked])], point_colors=[pal[t["year"]] for t in ranked], horizontal=True,
              maxv=max(sr(t, "P2AResponseRate") for t in ranked) * 1.18, fmt='0.00"%"', font=11, gap=40, label_size=12)
    if len(years) > 1:
        text(s, 0.5, 6.3, 7.5, 0.3, [[(f"{ylab(y)}   ", {"bold": True, "color": pal[y]}) for y in sorted(years)]], size=11, color=MUTED)
    rect(s, 8.4, 1.45, 4.4, 5.0, PANEL)
    text(s, 8.65, 1.6, 4.0, 4.8, panel_text(dk.get("ranking_panel", [])), size=12.5, space_after=7)
    source(s, SRC)

    refs = [r for r in (request.get("reference_templates") or []) if r["template"] in tpl]
    ref_notes = {r["template"]: r for r in (analysis.get("reference_templates") or [])}
    for i0 in range(0, len(refs), 3):
        chunk = refs[i0:i0 + 3]
        s = D.slide()
        ry = tpl[chunk[0]["template"]]["year"]
        header(s, ylab(ry).upper(), f"What the {ry} templates said")
        cw, gap, y0 = 3.95, 0.225, 1.5
        for i, r in enumerate(chunk):
            t = tpl[r["template"]]
            note = ref_notes.get(r["template"], {})
            x = 0.5 + i * (cw + gap)
            c = r.get("card", {})
            rect(s, x, y0, cw, 4.95, PANEL)
            rect(s, x, y0, cw, 0.5, NAVY if i < 2 else PURPLE)
            text(s, x + 0.2, y0, cw - 0.3, 0.5, f"{note.get('tag', r['template'])}  |  {M.nice_date(t['date'])}", size=14, bold=True, color=WHITE, anchor="m")
            text(s, x + 0.2, y0 + 0.58, cw - 0.4, 0.25, "Template: " + r["template"], size=10.5, color=MUTED, italic=True)
            text(s, x + 0.2, y0 + 0.85, 2.0, 0.55, M.fmt_pct(sr(t, "P2AResponseRate")), size=28, bold=True, color=RED, anchor="m")
            text(s, x + 2.15, y0 + 0.87, cw - 2.3, 0.55, ["click rate", "ReadRate " + M.fmt_pct(sr(t, "ReadRate")), "DeliveryRate " + M.fmt_pct(sr(t, "DeliveryRate"))], size=9.5, color=MUTED, anchor="m")
            rect(s, x + 0.2, y0 + 1.5, cw - 0.4, 0.015, "BFBFC6")
            paras = [[("Card Title", {"bold": True, "color": MUTED, "size": 10.5})], [(c.get("title", ""), {"bold": True, "color": NAVY, "size": 13})],
                     [("Card Description", {"bold": True, "color": MUTED, "size": 10.5})]]
            desc = c.get("description", "")
            for line in (desc if isinstance(desc, list) else [desc]):
                paras.append([(line, {"size": 11})])
            paras.append([("Button  ", {"bold": True, "color": MUTED, "size": 10.5}), (c.get("button", ""), {"bold": True, "color": RED, "size": 12})])
            if note.get("bullets"):
                paras.append([("What stands out", {"bold": True, "color": MUTED, "size": 10.5})])
                for bb in note["bullets"]:
                    paras.append({"runs": [(bb, {"size": 11})], "bullet": True})
            text(s, x + 0.2, y0 + 1.6, cw - 0.4, 3.3, paras, size=11, space_after=4)
        source(s, "Source: partner portal template export and text. No graphics available. Emojis may appear as placeholders.")

    s = D.slide()
    header(s, "ROOT CAUSE", "Where each hypothesis stands")
    cls = {"supported": (NAVY, WHITE), "directional": (NAVY, WHITE), "plausible": (YEL, TXT), "not_driver": ("D9D9E0", TXT), "unknown": ("D9D9E0", TXT)}
    rows = [["Hypothesis", "What the data shows", "Status"]]
    for h in analysis["hypotheses"][:8]:
        f, c = cls.get(h["status_class"], ("D9D9E0", TXT))
        rows.append([f"{h['id']}  {h['hypothesis']}", h["evidence"], (h["status"], {"fill": f, "color": c, "bold": True})])
    rh = min(0.68, 4.8 / len(rows))
    table(s, 0.5, 1.5, 12.3, [3.3, 6.5, 2.5], rows, size=11 if len(rows) <= 7 else 10, rowh=rh)
    text(s, 0.5, 6.4, 12.3, 0.4, dk.get("hypotheses_footnote", "Each template is a single send date. Audience, send time and post-click data (visits, orders) are not yet available."), size=12, italic=True, color=MUTED)
    source(s, SRC)

    def focus(n, t, a, rq):
        s = D.slide()
        header(s, f"TEMPLATE {n:02d}", f"Template in Focus  |  {rq.get('label') or t['template']}",
               sub=f"Template name: {t['template']}   |   Sent {M.nice_date(t['date'])}   |   GoogleID {t['google_id_short']}")
        img = os.path.join(image_dir, rq.get("image", "") or "")
        ph = 0
        if rq.get("image") and os.path.exists(img):
            _, ph = picture(s, img, 0.6, 1.75, 7.6, 3.1)
        my = 1.75 + ph + 0.2
        mh = 6.5 - my
        for i, (big, lab) in enumerate([(M.fmt_pct(sr(t, "DeliveryRate")), "DeliveryRate"), (M.fmt_pct(sr(t, "ReadRate")), "ReadRate"), (M.fmt_pct(sr(t, "P2AResponseRate")), "P2AResponseRate (click rate)")]):
            mx = 0.6 + i * 2.58
            rect(s, mx, my, 2.44, mh, PANEL)
            rect(s, mx, my, 0.06, mh, RED)
            text(s, mx + 0.2, my + 0.06, 2.2, mh * 0.55, big, size=22, bold=True, color=NAVY, anchor="m")
            text(s, mx + 0.2, my + mh * 0.58, 2.2, mh * 0.4, lab, size=10.5, color=MUTED)
        c = rq.get("card", {})
        desc = c.get("description", "")
        desc = desc if isinstance(desc, str) else "\n".join(desc)
        rect(s, 8.5, 1.75, 4.3, 4.75, PANEL)
        text(s, 8.75, 1.92, 3.8, 0.3, "CURRENT LIVE TEXT", size=14, bold=True, color=RED)
        text(s, 8.75, 2.3, 3.8, 0.25, "Card Title", size=12, bold=True, color=MUTED)
        text(s, 8.75, 2.55, 3.8, 0.75, c.get("title", ""), size=14, bold=True, color=NAVY)
        text(s, 8.75, 3.35, 3.8, 0.25, "Card Description", size=12, bold=True, color=MUTED)
        text(s, 8.75, 3.6, 3.8, 2.0, desc, size=12.5 if len(desc) < 260 else 11, color=TXT)
        text(s, 8.75, 5.65, 3.8, 0.25, "CTA", size=12, bold=True, color=MUTED)
        text(s, 8.75, 5.9, 3.8, 0.3, c.get("button", ""), size=12, bold=True, color=RED)
        source(s, "Source: partner portal template export; creative as supplied. Emojis may appear as placeholders.")
        notes(s, a.get("speaker_note", ""))

    def six(n, t, a, rq):
        s = D.slide()
        works, watch = a.get("works", []), a.get("watch_outs", [])
        nw = min(len(works), 3)
        items = [(x, NAVY) for x in works[:nw]] + [(x, RED) for x in watch[:6 - nw]]
        header(s, f"TEMPLATE {n:02d}  \u00B7  {rq.get('label') or t['template']}", "What Works and What to Watch" if works else "What Might Not Work in the Current Template")
        pos = [(0.5, 1.5), (0.5, 3.1), (0.5, 4.7), (6.85, 1.5), (6.85, 3.1), (6.85, 4.7)]
        for i, ((x, y), (it, col)) in enumerate(zip(pos, items)):
            rect(s, x, y + 0.03, 0.06, 1.35, col)
            text(s, x + 0.22, y, 5.7, 0.4, f"{i + 1}.  {it['h']}", size=16, bold=True, color=NAVY)
            text(s, x + 0.22, y + 0.42, 5.7, 1.0, it["b"], size=13.5 if len(it["b"]) < 170 else 12, color=TXT)
        rect(s, 0.5, 6.2, 12.3, 0.4, PANEL)
        text(s, 0.7, 6.2, 11.9, 0.4, a.get("strip") or f"Click rate {M.fmt_pct(sr(t, 'P2AResponseRate'))}", size=12, color=NAVY, bold=True, anchor="m")

    def variations(n, t, a, rq):
        vs = a.get("variations") or []
        if not vs:
            return
        s = D.slide()
        header(s, f"TEMPLATE {n:02d}  \u00B7  {rq.get('label') or t['template']}", "Template Recommendations")
        for (x, w, fill, v) in [(0.4, 6.05, NAVY, vs[0]), (6.85, 5.95, PURPLE, vs[1] if len(vs) > 1 else None)]:
            if not v:
                continue
            rect(s, x, 1.5, w, 5.05, PANEL)
            rect(s, x, 1.5, w, 0.5, fill)
            text(s, x + 0.3, 1.5, w - 0.5, 0.5, v["name"], size=16, bold=True, color=WHITE, anchor="m")
            text(s, x + 0.3, 2.2, w - 0.6, 0.25, "Card Title", size=12, bold=True, color=MUTED)
            text(s, x + 0.3, 2.45, w - 0.6, 0.9, v["title"], size=20, bold=True, color=NAVY)
            text(s, x + 0.3, 3.4, w - 0.6, 0.25, "Card Description", size=12, bold=True, color=MUTED)
            text(s, x + 0.3, 3.65, w - 0.6, 1.5, v["description"] if isinstance(v["description"], list) else [v["description"]], size=15, color=TXT, space_after=6)
            text(s, x + 0.3, 5.45, w - 0.6, 0.25, "Why test it", size=12, bold=True, color=MUTED)
            text(s, x + 0.3, 5.7, w - 0.6, 0.8, v["why"], size=13, color=TXT)

    for n, t in enumerate([t for t in metrics["templates"] if t["year"] == cur], 1):
        a, rq = ana_t[t["template"]], req_t[t["template"]]
        focus(n, t, a, rq)
        six(n, t, a, rq)
        variations(n, t, a, rq)

    s = D.slide()
    header(s, "RECOMMENDATIONS", dk.get("recommendations_title", "Changes to make in the next sends"))
    recs = dk.get("recommendations") or []
    cw, ch, gx, gy, y0 = 3.95, 2.35, 0.225, 0.25, 1.5
    for i, r in enumerate(recs[:6]):
        x, y = 0.5 + (i % 3) * (cw + gx), y0 + (i // 3) * (ch + gy)
        rect(s, x, y, cw, ch, PANEL)
        rect(s, x, y, 0.06, ch, RED)
        rect(s, x + 0.25, y + 0.2, 0.5, 0.5, NAVY, shape=MSO_SHAPE.OVAL)
        text(s, x + 0.25, y + 0.2, 0.5, 0.5, str(i + 1), size=16, bold=True, color=WHITE, align="c", anchor="m")
        text(s, x + 0.9, y + 0.2, cw - 1.1, 0.7, r["h"], size=14, bold=True, color=NAVY, anchor="m")
        text(s, x + 0.25, y + 1.0, cw - 0.45, 1.3, r["b"], size=13 if len(r["b"]) > 120 else 13.5, color=TXT)
    source(s, "Source: template review; partner portal template export.")

    s = D.slide()
    asks = dk.get("data_asks") or []
    header(s, "NEXT STEPS", f"One test and {len(asks[:3])} data ask" + ("s" if len(asks[:3]) != 1 else ""))
    rect(s, 0.5, 1.5, 4.7, 5.0, NAVY)
    text(s, 0.8, 1.7, 4.2, 0.4, "Test in the next window", size=18, bold=True, color=WHITE)
    text(s, 0.8, 2.3, 4.2, 4.0, dk.get("next_steps_test", ["Run one clean A/B test", "Add a tracking tag to every button URL"]), size=13.5, color=WHITE, bullets=True, space_after=10)
    for i, a_ in enumerate(asks[:3]):
        y = 1.5 + i * 1.7
        rect(s, 5.5, y, 7.3, 1.55, PANEL)
        rect(s, 5.5, y, 0.07, 1.55, RED)
        text(s, 5.8, y + 0.1, 6.8, 0.35, a_["owner"], size=16, bold=True, color=NAVY)
        text(s, 5.8, y + 0.5, 6.8, 1.0, a_["text"], size=13, color=TXT)
    source(s, "Source: template review.")

    D.thanks_slide()
    for sl in D.prs.slides:                       # python-pptx writes negative axis ids; make them valid unsigned ints
        for shp in sl.shapes:
            if shp.has_chart:
                for el in shp.chart._chartSpace.iter():
                    if el.tag.endswith("}axId") or el.tag.endswith("}crossAx"):
                        el.set("val", str(abs(int(el.get("val")))))
    D.prs.save(out_path)
    return out_path
