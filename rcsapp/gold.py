"""
Gold set: human content ratings used to measure how good the engines are. It is NOT training data.

The rule that makes it valid: a template is rated from its card text and image BEFORE its results are seen.
Rows rated afterwards are marked saw_performance = Y and are left out of the benchmark.

What it measures
  * Content risk (Low / Medium / High): agreement between the engine and the human, as quadratic weighted kappa,
    shown beside human-vs-human agreement where a second rater scored the same templates (the ceiling).
  * Text-check flags: per-check recall and precision against the human's flags.
  * Outcomes (secondary): rank correlation between the engine's content score and click index, inside each brand.
    Reported only; outcomes are confounded by audience and timing.
"""
import hashlib
import io

import numpy as np
import pandas as pd

from .engines import ml_engine as E

RISKS = ["Low", "Medium", "High"]
SEV = {r: i for i, r in enumerate(RISKS)}
REC_THEMES = ["name_offer", "add_number", "one_action", "front_load", "remove_chance", "shorten", "fewer_variables", "tone_formal", "cta_low_commitment", "other"]
BARS = {"assist": 0.5, "light": 0.7, "flag_recall": 0.7, "flag_precision": 0.6, "stability": 0.9, "min_positives": 5, "min_rows_first_read": 30, "min_rows_per_rule": 60}

RUBRIC = [
    ("Low", "The reader can tell what they get and what to do, and nothing plausibly costs clicks."),
    ("Medium", "The offer is clear, but one or two weaknesses likely cost clicks."),
    ("High", "The reader cannot tell what they get, or the image and the text conflict, or conditions are hidden or contradictory, or there are three or more weaknesses."),
]


def flag_cols(rules):
    return [r["id"] for r in rules["watch"]]


def gold_cols(rules):
    return (["gold_id", "brand", "vertical", "template_id", "title", "description", "button_label", "creative_note", "image_file", "rater", "rated_on",
             "saw_performance", "content_risk", "confidence"] + flag_cols(rules) +
            ["other_issue", "rec_theme_1", "rec_theme_2", "rec_theme_3", "second_rater", "second_risk", "split", "outcome_click_index"])


def assign_split(df, test_share=0.3, seed="v1"):
    """All templates of a brand go to the same split where possible (so a brand's style is not on both sides)."""
    df = df.copy()
    brands = sorted(df["brand"].unique(), key=lambda b: hashlib.md5((str(b) + seed).encode()).hexdigest())
    if len(brands) < 2:
        h = df["template_id"].map(lambda t: int(hashlib.md5((str(t) + seed).encode()).hexdigest(), 16) % 100)
        df["split"] = np.where(h < int(test_share * 100), "test", "dev")
        return df
    total, acc, test = len(df), 0, set()
    for b in brands:
        if acc / total >= test_share:
            break
        test.add(b)
        acc += int((df["brand"] == b).sum())
    df["split"] = np.where(df["brand"].isin(test), "test", "dev")
    return df


def blind_sheet(hist, rules, seed=0):
    """Rating sheet built from the training history: one row per text, results removed, order shuffled."""
    u = hist[hist["reason"].ne("excluded")].copy()
    u = u[(u["title"].astype(str).str.strip() != "") | (u["description"].astype(str).str.strip() != "")]
    u = u.sort_values(["brand", "date"]).groupby("text_group", as_index=False).first()
    u = u.sample(frac=1.0, random_state=seed).reset_index(drop=True)
    cols = gold_cols(rules)
    g = pd.DataFrame({c: "" for c in cols}, index=range(len(u)))
    g["gold_id"] = [f"G{i + 1:03d}" for i in range(len(u))]
    g["brand"], g["vertical"], g["template_id"] = u["brand"], u["vertical"], u["template"]
    g["title"], g["description"], g["button_label"] = u["title"], u["description"], u["button"]
    g["saw_performance"] = "N"
    g = assign_split(g)
    return g[cols]


def flag_cols_from(df):
    known = ["no_benefit_named", "not_quantified", "chance_based", "urgency_no_benefit", "floor_or_ceiling", "long_description",
             "many_variables", "shouty", "competing_actions", "thin_title"]
    return [c for c in df.columns if c in known]


def template_workbook(df, rules):
    """An .xlsx the rater fills in, with drop-downs, a rubric and instructions."""
    from openpyxl.worksheet.datavalidation import DataValidation
    cols = list(df.columns)
    bio = io.BytesIO()
    with pd.ExcelWriter(bio, engine="openpyxl") as w:
        df.to_excel(w, sheet_name="Rate here", index=False)
        pd.DataFrame({"How to rate": [
            "1. Rate each row from its title, description, button and image BEFORE you open any performance figure. This sheet has none on purpose.",
            "2. content_risk: Low, Medium or High, using the Rubric sheet. confidence: High, Medium or Low (Low rows are left out of the headline score).",
            "3. Put Y in each flag column that applies. Use other_issue for anything the flags miss; new checks come from there.",
            "4. rec_theme_1 to 3: what you would change first. Pick from the list in the Rubric sheet.",
            "5. If you have already seen a template's results, set saw_performance to Y. It stays in the sheet for study but is left out of the benchmark.",
            "6. A second rater fills second_rater and second_risk on at least 15 rows. That gives the human-versus-human ceiling.",
            "7. Do not edit gold_id, split or outcome_click_index."]}).to_excel(w, sheet_name="How to", index=False)
        pd.DataFrame({"Risk": [r[0] for r in RUBRIC], "Meaning": [r[1] for r in RUBRIC]}).to_excel(w, sheet_name="Rubric", index=False, startrow=0)
        pd.DataFrame({"rec_theme": REC_THEMES}).to_excel(w, sheet_name="Rubric", index=False, startcol=3)
        ws = w.sheets["Rate here"]
        n = len(df) + 1
        for name, options in (("content_risk", RISKS), ("second_risk", RISKS), ("saw_performance", ["N", "Y"]), ("confidence", ["High", "Medium", "Low"]),
                              ("rec_theme_1", REC_THEMES), ("rec_theme_2", REC_THEMES), ("rec_theme_3", REC_THEMES)):
            if name in cols:
                col = cols.index(name) + 1
                dv = DataValidation(type="list", formula1='"' + ",".join(options) + '"', allow_blank=True)
                ws.add_data_validation(dv)
                dv.add(f"{ws.cell(row=2, column=col).coordinate}:{ws.cell(row=max(n, 2), column=col).coordinate}")
        for fc in flag_cols_from(df):
            col = cols.index(fc) + 1
            dv = DataValidation(type="list", formula1='"Y"', allow_blank=True)
            ws.add_data_validation(dv)
            dv.add(f"{ws.cell(row=2, column=col).coordinate}:{ws.cell(row=max(n, 2), column=col).coordinate}")
        ws.freeze_panes = "E2"
        for c, width in (("A", 8), ("B", 16), ("D", 24), ("E", 40), ("F", 70)):
            ws.column_dimensions[c].width = width
    return bio.getvalue()


def read_gold(data):
    df = pd.read_excel(io.BytesIO(data), sheet_name=0, dtype=object) if isinstance(data, (bytes, bytearray)) else pd.read_excel(data, sheet_name=0, dtype=object)
    for c in ("gold_id", "brand", "template_id", "title", "description"):
        if c not in df.columns:
            raise ValueError(f"This is not a gold-set sheet made by this app. Missing column: {c}.")
    df = df[df["gold_id"].notna()].copy()
    for c in df.columns:
        df[c] = df[c].map(lambda v: "" if v is None or (isinstance(v, float) and np.isnan(v)) else str(v).strip())
    return df


def _qwk(a, b):
    from sklearn.metrics import cohen_kappa_score
    if len(a) < 2:
        return None
    try:
        k = cohen_kappa_score(a, b, weights="quadratic", labels=[0, 1, 2])
    except Exception:
        return None
    return None if (k is None or np.isnan(k)) else float(k)


def engine_scores(gold, rules):
    rows = []
    for r in gold.itertuples():
        e = E.evaluate({"title": r.title, "description": r.description, "button": getattr(r, "button_label", "")}, rules)
        rows.append({"gold_id": r.gold_id, "engine_risk": e["risk"], "engine_score": e["score"], "engine_flags": {w["id"] for w in e["watch"]}})
    return pd.DataFrame(rows)


def score(gold, rules, hist=None, llm_risk=None):
    """llm_risk: optional {template_id: 'Low'|'Medium'|'High'} from the language model. Returns a dict of results."""
    out = {"notes": []}
    g = gold.copy()
    g = g.merge(engine_scores(g, rules), on="gold_id", how="left")
    if llm_risk:
        g["llm_risk"] = g["template_id"].map(llm_risk)
    rated = g[g["content_risk"].isin(RISKS)]
    blind = rated[rated["saw_performance"].str.upper().ne("Y")]
    head = blind[blind["confidence"].str.title().ne("Low")]
    out["counts"] = {"rows": int(len(g)), "rated": int(len(rated)), "blind": int(len(blind)), "headline": int(len(head)),
                     "excluded_saw_performance": int((rated["saw_performance"].str.upper() == "Y").sum()),
                     "excluded_low_confidence": int((blind["confidence"].str.title() == "Low").sum())}
    if len(head) < BARS["min_rows_first_read"]:
        out["notes"].append(f"Only {len(head)} blind, confident ratings. {BARS['min_rows_first_read']} are the minimum for a first read and the margin of error is wide below {BARS['min_rows_per_rule']}.")

    def block(df, label):
        a = df["content_risk"].map(SEV).astype(int)
        b = df["engine_risk"].map(SEV).astype(int)
        return {"set": label, "n": int(len(df)), "qwk": _qwk(a, b), "exact": float((a == b).mean()) if len(df) else None}

    res = [block(head, "all")] if len(head) else []
    for sp in ("dev", "test"):
        sub = head[head["split"] == sp]
        if len(sub):
            res.append(block(sub, sp))
    out["engine_vs_human"] = res
    conf = pd.crosstab(head["content_risk"], head["engine_risk"]).reindex(index=RISKS, columns=RISKS, fill_value=0) if len(head) else pd.DataFrame()
    out["confusion"] = conf
    if llm_risk:
        h2 = head[head["llm_risk"].isin(RISKS)]
        if len(h2):
            out["llm_vs_human"] = {"n": int(len(h2)), "qwk": _qwk(h2["content_risk"].map(SEV).astype(int), h2["llm_risk"].map(SEV).astype(int)),
                                   "exact": float((h2["content_risk"] == h2["llm_risk"]).mean())}
    sec = blind[blind["second_risk"].isin(RISKS)]
    out["human_vs_human"] = {"n": int(len(sec)), "qwk": _qwk(sec["content_risk"].map(SEV).astype(int), sec["second_risk"].map(SEV).astype(int)) if len(sec) else None}
    if len(sec) < 15:
        out["notes"].append(f"Only {len(sec)} rows have a second rater. At least 15 are needed to know how well two humans agree, which is the ceiling for any engine.")

    fl = []
    for fc in flag_cols(rules):
        if fc not in head.columns:
            continue
        human = head[fc].str.upper().eq("Y")
        eng = head["engine_flags"].map(lambda s, f=fc: f in s)
        tp, fp, fn = int((human & eng).sum()), int((~human & eng).sum()), int((human & ~eng).sum())
        pos = int(human.sum())
        rec = tp / (tp + fn) if (tp + fn) else None
        prec = tp / (tp + fp) if (tp + fp) else None
        if pos < BARS["min_positives"]:
            verdict = f"too few human positives ({pos}) to judge"
        elif rec is not None and prec is not None and rec >= BARS["flag_recall"] and prec >= BARS["flag_precision"]:
            verdict = "keep"
        else:
            verdict = "rewrite or drop"
        fl.append({"check": fc, "human_positives": pos, "engine_fires": int(eng.sum()), "recall": None if rec is None else round(rec, 2), "precision": None if prec is None else round(prec, 2), "verdict": verdict})
    out["flags"] = pd.DataFrame(fl)

    if hist is not None and len(hist):
        idx = hist[hist["usable"]].groupby(["brand", "template"])["click_index"].mean().reset_index().rename(columns={"template": "template_id", "click_index": "outcome"})
        m = g.merge(idx, on=["brand", "template_id"], how="inner")
        cors = []
        for b, sub in m.groupby("brand"):
            if len(sub) >= 5 and sub["engine_score"].nunique() > 1:
                cors.append(float(sub["engine_score"].corr(sub["outcome"], method="spearman")))
        out["outcomes"] = {"templates_matched": int(len(m)), "brands_with_5_or_more": len(cors), "mean_spearman": (float(np.mean(cors)) if cors else None),
                           "note": "A negative value is the expected direction: a higher risk score goes with a lower click index. Outcomes are confounded by audience and timing, so this is secondary."}

    q = out["engine_vs_human"][0]["qwk"] if out["engine_vs_human"] else None
    if q is None:
        out["verdict"] = "Not enough rated rows to say."
    elif q >= BARS["light"]:
        out["verdict"] = f"Agreement is {q:.2f}: close enough to lighten the review. Read the recommendations and anything about phone behaviour."
    elif q >= BARS["assist"]:
        out["verdict"] = f"Agreement is {q:.2f}: use the engine to assist, and still read every content rating line by line."
    else:
        out["verdict"] = f"Agreement is {q:.2f}, below the {BARS['assist']} bar: do not rely on the engine's content ratings; revise the checks first."
    return out


def scorecard_markdown(res):
    L = ["# Gold-set scorecard", f"Verdict: {res['verdict']}", "", f"Counts: {res['counts']}", ""]
    L.append("## Engine vs human content risk")
    for r in res["engine_vs_human"]:
        L.append(f"- {r['set']}: n={r['n']}, weighted kappa={r['qwk'] if r['qwk'] is None else round(r['qwk'], 2)}, exact agreement={None if r['exact'] is None else round(r['exact'], 2)}")
    if "llm_vs_human" in res:
        L.append(f"- language model: {res['llm_vs_human']}")
    L.append(f"- human vs human: {res['human_vs_human']}")
    L.append("")
    if len(res.get("flags", [])):
        L.append("## Text-check flags\n" + res["flags"].to_string(index=False))
    if "outcomes" in res:
        L.append(f"\n## Outcomes (secondary)\n{res['outcomes']}")
    L += ["", "## Notes"] + [f"- {n}" for n in res["notes"]]
    return "\n".join(L)
