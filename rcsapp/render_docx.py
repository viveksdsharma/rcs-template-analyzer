"""Word report builder (python-docx). Layout follows the approved report format. Footer: page number only."""
import os

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from . import metrics as M
from .features import iconify

FILL = {"green": "E6F4EA", "amber": "FFF4D6", "red": "FDE9E7", "grey": "F2F2F2"}
LINE = "BFBFBF"
PAGE_W = 9746


def _hex(s):
    return RGBColor.from_string(s)


def _fill_of(x):
    return FILL.get(x, x)


def rich(p, text, size=9.5, color=None, bold=False, italic=False):
    for i, part in enumerate(str(text).split("**")):
        if not part:
            continue
        r = p.add_run(part)
        r.font.size = Pt(size)
        r.bold = bold or (i % 2 == 1)
        r.italic = italic
        if color:
            r.font.color.rgb = _hex(color)
    return p


def shade(cell, hexfill):
    tcPr = cell._tc.get_or_add_tcPr()
    for e in tcPr.findall(qn("w:shd")):
        tcPr.remove(e)
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hexfill)
    tcPr.append(shd)


def _tbl_setup(table, widths):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tblPr = table._tbl.tblPr
    lay = OxmlElement("w:tblLayout")
    lay.set(qn("w:type"), "fixed")
    tblPr.append(lay)
    tw = tblPr.find(qn("w:tblW"))
    if tw is None:
        tw = OxmlElement("w:tblW")
        tblPr.append(tw)
    tw.set(qn("w:type"), "dxa")
    tw.set(qn("w:w"), str(sum(widths)))
    borders = OxmlElement("w:tblBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        b = OxmlElement(f"w:{side}")
        b.set(qn("w:val"), "single")
        b.set(qn("w:sz"), "4")
        b.set(qn("w:space"), "0")
        b.set(qn("w:color"), LINE)
        borders.append(b)
    tblPr.append(borders)
    mar = OxmlElement("w:tblCellMar")
    for side, v in (("top", 40), ("left", 90), ("bottom", 40), ("right", 90)):
        e = OxmlElement(f"w:{side}")
        e.set(qn("w:w"), str(v))
        e.set(qn("w:type"), "dxa")
        mar.append(e)
    tblPr.append(mar)
    for i, gc in enumerate(table._tbl.tblGrid.findall(qn("w:gridCol"))):
        gc.set(qn("w:w"), str(widths[i]))


def _row_flags(row, header=False):
    trPr = row._tr.get_or_add_trPr()
    trPr.append(OxmlElement("w:cantSplit"))
    if header:
        trPr.append(OxmlElement("w:tblHeader"))


def write_cell(cell, content, size=8.5, bold=False, color=None, fill=None, width=None):
    if width:
        cell.width = Inches(width / 1440)
    if fill:
        shade(cell, _fill_of(fill))
    items = content if isinstance(content, list) else [content]
    first = True
    for it in items:
        p = cell.paragraphs[0] if first else cell.add_paragraph()
        first = False
        p.paragraph_format.space_after = Pt(1)
        p.paragraph_format.space_before = Pt(0)
        if isinstance(it, dict) and it.get("bullet"):
            p.paragraph_format.left_indent = Inches(0.14)
            p.paragraph_format.first_line_indent = Inches(-0.14)
            rich(p, "\u2022 " + it["bullet"], size=size, color=color)
        elif isinstance(it, dict) and it.get("picture"):
            p.add_run().add_picture(it["picture"], width=Inches(it.get("width", 2.4)))
        else:
            rich(p, it, size=size, bold=bold, color=color)


def add_table(doc, widths, header, rows, size=8.5, hsize=None, header_fill="CC0000"):
    t = doc.add_table(rows=0, cols=len(widths))
    _tbl_setup(t, widths)
    if header:
        r = t.add_row()
        _row_flags(r, header=True)
        for i, h in enumerate(header):
            write_cell(r.cells[i], h, size=hsize or size, bold=True, color="FFFFFF", fill=header_fill, width=widths[i])
    for row in rows:
        r = t.add_row()
        _row_flags(r)
        for i, c in enumerate(row):
            if isinstance(c, dict) and "t" in c:
                write_cell(r.cells[i], c["t"], size=size, bold=c.get("bold", False), fill=c.get("fill"), width=widths[i])
            else:
                write_cell(r.cells[i], c, size=size, width=widths[i])
    return t


def ratio_widths(*r):
    s = sum(r)
    w = [int(PAGE_W * x / s) for x in r]
    w[0] += PAGE_W - sum(w)
    return w


class Doc:
    def __init__(self, font, heading_color):
        self.d = Document()
        sec = self.d.sections[0]
        sec.page_width, sec.page_height = Inches(8.27), Inches(11.69)
        sec.left_margin = sec.right_margin = Inches(0.75)
        sec.top_margin, sec.bottom_margin = Inches(0.6), Inches(0.55)
        st = self.d.styles["Normal"]
        st.font.name = font
        st.font.size = Pt(10)
        st.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), font)
        for name, size in (("Heading 1", 13), ("Heading 2", 10.5)):
            h = self.d.styles[name]
            h.font.name = font
            h.font.size = Pt(size)
            h.font.bold = True
            h.font.color.rgb = _hex(heading_color)
            rf = h.element.get_or_add_rPr().get_or_add_rFonts()
            for att in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"):
                rf.attrib.pop(qn("w:" + att), None)
            for att in ("ascii", "hAnsi", "eastAsia"):
                rf.set(qn("w:" + att), font)
            h.paragraph_format.keep_with_next = True
            h.paragraph_format.space_before = Pt(13 if name == "Heading 1" else 9)
            h.paragraph_format.space_after = Pt(5 if name == "Heading 1" else 3.5)
        fp = sec.footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        r = fp.add_run("Page ")
        r.font.size = Pt(7.5)
        r.font.color.rgb = _hex("7F7F7F")
        run = fp.add_run()
        run.font.size = Pt(7.5)
        run.font.color.rgb = _hex("7F7F7F")
        for kind, txt in (("begin", None), (None, "PAGE"), ("end", None)):
            if kind:
                fc = OxmlElement("w:fldChar")
                fc.set(qn("w:fldCharType"), kind)
                run._r.append(fc)
            else:
                it = OxmlElement("w:instrText")
                it.set(qn("xml:space"), "preserve")
                it.text = txt
                run._r.append(it)

    def h1(self, t):
        self.d.add_heading(t, level=1)

    def h2(self, t):
        self.d.add_heading(t, level=2)

    def p(self, text, size=9.5, after=4.5, keep=False, italic=False, color=None, before=0):
        p = self.d.add_paragraph()
        p.paragraph_format.space_after = Pt(after)
        p.paragraph_format.space_before = Pt(before)
        p.paragraph_format.keep_with_next = keep
        rich(p, text, size=size, italic=italic, color=color)
        return p

    def note(self, text):
        return self.p(text, size=7.5, italic=True, color="595959", after=4.5, before=1.5)

    def bullet(self, text, size=9.5):
        p = self.d.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.25)
        p.paragraph_format.first_line_indent = Inches(-0.17)
        p.paragraph_format.space_after = Pt(2.5)
        rich(p, "\u2022  " + text, size=size)

    def numbered(self, items, size=9.5):
        for i, text in enumerate(items, 1):
            p = self.d.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.27)
            p.paragraph_format.first_line_indent = Inches(-0.22)
            p.paragraph_format.space_after = Pt(2.5)
            rich(p, f"{i}.  " + text, size=size)

    def gap(self):
        self.d.add_paragraph().paragraph_format.space_after = Pt(2.5)

    def save(self, path):
        self.d.save(path)


def _hb(items):
    return [f"**{x['h']}.** {x['b']}" if x.get("h") else x["b"] for x in items]


def _risk_fill(r):
    return {"Low": "green", "Medium": "amber", "High": "red"}.get(r, "grey")


def _prio_fill(p):
    return {"Now": "green", "Next": "amber"}.get(p, "grey")


def _status_fill(c):
    return {"supported": "red", "directional": "red", "plausible": "amber", "not_driver": "grey", "unknown": "grey"}.get(c, "grey")


CHANGE_STYLE = {"DeliveryRate": "pts", "ReadRate": "pts", "pre_send_failure_rate": "pts", "end_to_end_delivery_rate": "pts",
                "P2AResponseRate": "fold", "expiry_rate": "fold", "click_rate_delivered_base": "pct"}
YOY_DEC = {"pre_send_failure_rate": 1, "end_to_end_delivery_rate": 1}
YOY_ROWS = [("P2AResponseRate", "P2AResponseRate (click rate)"), ("ReadRate", "ReadRate"), ("DeliveryRate", "DeliveryRate"),
            ("expiry_rate", "Expiry rate*"), ("pre_send_failure_rate", "Pre-send failure rate*"),
            ("end_to_end_delivery_rate", "End-to-end delivery rate*"), ("click_rate_delivered_base", "Click rate, delivered base*")]


def _change_text(key, m):
    style = CHANGE_STYLE.get(key, "pts")
    if style == "fold" and m["fold"]:
        return f"{1 / m['fold']:.1f}x lower" if m["fold"] < 1 else f"{m['fold']:.1f}x higher"
    if style == "pct" and m["fold"]:
        return f"{(m['fold'] - 1) * 100:+.0f}%"
    return f"{m['change_pts']:+.1f} points"


def build_docx(metrics, analysis, request, out_path, style, image_dir):
    font = style.get("docx_font", "Times New Roman")
    hc = style.get("heading_color", "C00000")
    hf = style.get("table_header_fill", "CC0000")
    D = Doc(font, hc)
    brand = request.get("brand_display") or request.get("brand")
    labels = request.get("period_labels", {})
    years = metrics["meta"]["years"]
    cur = max(years)
    ylabel = lambda y: labels.get(str(y), str(y))  # noqa: E731
    tpl = {t["template"]: t for t in metrics["templates"]}
    req_t = {t["template"]: t for t in request["templates"]}
    ana_t = {t["template"]: t for t in analysis["templates"]}
    n = 0

    def nxt():
        nonlocal n
        n += 1
        return n

    p = D.d.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    rich(p, analysis["meta"].get("title") or f"{brand}: Template Analysis", size=20, bold=True, color=hc)
    if analysis["meta"].get("subtitle"):
        D.p(analysis["meta"]["subtitle"], size=11, color="404040", after=6)

    D.h1(f"{nxt()}. Summary")
    D.p(analysis["summary"]["lead"], size=9.5)
    for b in analysis["summary"]["bullets"]:
        D.bullet(b)

    ext_tables = request.get("brand_reported_tables") or analysis.get("external_tables")
    if ext_tables:
        D.h1(f"{nxt()}. Current Status")
        for k, tb in enumerate(ext_tables, 1):
            D.h2(f"{n}.{k} {tb['title']}")
            if tb.get("intro"):
                D.p(tb["intro"], keep=True)
            add_table(D.d, ratio_widths(*([1.6] + [1] * (len(tb["header"]) - 1))), tb["header"], tb["rows"], size=8.5, header_fill=hf)
            if tb.get("note"):
                D.note(tb["note"])

    sec3 = nxt()
    D.h1(f"{sec3}. Template data")
    D.p("Source: the template export from the partner portal. Each row is one Date. Column names below are the sheet's own. "
        "P2AResponseRate is the click rate (P2AResponses \u00F7 Messages Read). Ratios the sheet does not contain are marked * and named for what they are used for.")
    D.h2(f"{sec3}.1 Counts as in the sheet")
    w31 = [2150, 780, 780, 900, 850, 900, 780, 830, 790, 986]
    hdr31 = ["Template", "Date", "GoogleID", "Messages Submitted", "Messages Sent", "Messages Delivered", "Messages Read", "Messages Failed", "Messages Expired", "P2AResponses"]
    cnt = ("submitted", "sent", "delivered", "read", "failed", "expired", "p2a_responses")
    for y in sorted(years, reverse=True):
        ts = [t for t in metrics["templates"] if t["year"] == y]
        D.p(f"**{ylabel(y)}**", keep=True, after=2)
        rows = [[t["template"], M.nice_date(t["date"]).rsplit(" ", 1)[0], t["google_id_short"]] + [M.fmt_int(t["counts"][k]) for k in cnt] for t in ts]
        tot = metrics["groups"]["year"][str(y)]["counts"]
        rows.append([{"t": f"**Total ({len(ts)} sends)**", "fill": "grey"}, {"t": "", "fill": "grey"}, {"t": "", "fill": "grey"}] + [{"t": f"**{M.fmt_int(tot[k])}**", "fill": "grey"} for k in cnt])
        add_table(D.d, w31, hdr31, rows, size=7, hsize=6.5, header_fill=hf)
        D.gap()
    D.note('GoogleID is shown without the brand prefix and the "_agent" suffix.')

    D.h2(f"{sec3}.2 Rates")
    D.p("DeliveryRate, ReadRate and P2AResponseRate are the sheet's own columns (P2AResponseRate is the click rate). The rates marked * are our calculations from the sheet counts, defined below the table.")
    w32 = [2250, 960, 740, 1150, 860, 880, 720, 880, 1306]
    hdr32 = ["Template", "DeliveryRate", "ReadRate", "P2AResponseRate (click rate)", "Pre-send failure rate*", "End-to-end delivery rate*", "Expiry rate*", "Read rate, sent base*", "Click rate, delivered base*"]
    for y in sorted(years, reverse=True):
        D.p(f"**{ylabel(y)}**", keep=True, after=2)
        rows = []
        for t in [t for t in metrics["templates"] if t["year"] == y]:
            r = t["rates"]
            rows.append([t["template"], M.fmt_pct(M.shown_rate(t, "DeliveryRate")), M.fmt_pct(M.shown_rate(t, "ReadRate")), "**" + M.fmt_pct(M.shown_rate(t, "P2AResponseRate")) + "**",
                         M.fmt_pct(r["pre_send_failure_rate"], 1), M.fmt_pct(r["end_to_end_delivery_rate"], 1), M.fmt_pct(r["expiry_rate"], 1), M.fmt_pct(r["read_rate_sent_base"]), M.fmt_pct(r["click_rate_delivered_base"])])
        g = metrics["groups"]["year"][str(y)]["rates"]
        rows.append([{"t": f"**Total, {y}**", "fill": "grey"}] + [{"t": "**" + x + "**", "fill": "grey"} for x in (
            M.fmt_pct(g["DeliveryRate"]), M.fmt_pct(g["ReadRate"]), M.fmt_pct(g["P2AResponseRate"]), M.fmt_pct(g["pre_send_failure_rate"], 1),
            M.fmt_pct(g["end_to_end_delivery_rate"], 1), M.fmt_pct(g["expiry_rate"], 1), M.fmt_pct(g["read_rate_sent_base"]), M.fmt_pct(g["click_rate_delivered_base"]))])
        add_table(D.d, w32, hdr32, rows, size=7, hsize=6.5, header_fill=hf)
        D.gap()
    add_table(D.d, [2300, 3300, 4146], ["Name", "Calculation", "Used for"], [[d["name"], d["formula"].replace(" / ", " \u00F7 "), d["used_for"]] for d in metrics["definitions"]], size=7.5, header_fill=hf)
    D.note("Totals and * rates are rounded. The sheet truncates its template-level rates to two decimals, so a value can differ by 0.01 point from a rounded recalculation.")

    D.h2(f"{sec3}.3 By GoogleID")
    rows = []
    for key, g in sorted(metrics["groups"]["year|google_id"].items(), key=lambda kv: kv[0].split("|")[0], reverse=True):
        y, gid = key.split("|", 1)
        names = ", ".join(dict.fromkeys(t["template"] for t in metrics["templates"] if str(t["year"]) == y and t["google_id"] == gid))
        r = g["rates"]
        rows.append([M.short_agent_id(gid), y, names, M.fmt_int(g["counts"]["submitted"]), M.fmt_pct(r["pre_send_failure_rate"], 1), M.fmt_pct(r["DeliveryRate"]), M.fmt_pct(r["ReadRate"]), M.fmt_pct(r["P2AResponseRate"])])
    add_table(D.d, [1000, 600, 2500, 1150, 1100, 1050, 850, 1496], ["GoogleID", "Year", "Templates", "Messages Submitted", "Pre-send failure rate*", "DeliveryRate", "ReadRate", "P2AResponseRate (click rate)"], rows, size=7.5, header_fill=hf)
    D.note("Rates by GoogleID are recomputed from the summed counts.")

    if analysis.get("numbers_commentary"):
        D.h2(f"{sec3}.4 What the numbers show ({ylabel(cur)})")
        D.numbered(_hb(analysis["numbers_commentary"]))

    yoy = metrics["comparisons"].get("year_over_year")
    if yoy:
        D.h2(f"{sec3}.5 {ylabel(yoy['from_year'])} vs {ylabel(yoy['to_year'])}")
        D.p("Totals across the templates of each period, recomputed from summed counts.", keep=True)
        ya, yb = metrics["groups"]["year"][str(yoy["from_year"])], metrics["groups"]["year"][str(yoy["to_year"])]
        rows = [["Messages Submitted", M.fmt_int(ya["counts"]["submitted"]), M.fmt_int(yb["counts"]["submitted"]), f"{yoy['submitted_growth'] * 100:+.0f}%"]]
        for key, label in YOY_ROWS:
            m = yoy["metrics"].get(key)
            if not m:
                continue
            hot = abs(m["change_pts"]) > 5 or (m["fold"] and (m["fold"] > 1.5 or m["fold"] < 0.67))
            f = "red" if hot else None
            d_ = YOY_DEC.get(key, 2)
            rows.append([{"t": f"**{label}**" if hot else label, "fill": f}, {"t": M.fmt_pct(m["from"], d_), "fill": f}, {"t": M.fmt_pct(m["to"], d_), "fill": f},
                         {"t": f"**{_change_text(key, m)}**" if hot else _change_text(key, m), "fill": f}])
        add_table(D.d, [3200, 2000, 2000, 2546], ["Metric", f"{ylabel(yoy['from_year'])} ({ya['n_templates']} sends)", f"{ylabel(yoy['to_year'])} ({yb['n_templates']} sends)", "Change"], rows, size=8.5, header_fill=hf)
        D.gap()
        if analysis.get("yoy_commentary"):
            D.numbered(_hb(analysis["yoy_commentary"]))

    sec4 = nxt()
    D.h1(f"{sec4}. Creative and text review")
    D.p("Each card combines a media image, a title, a description and one URL button. Content risk ratings are our view of the content, made from the content alone and shown beside the results for comparison. Emojis in card text may appear as placeholders.")
    cur_templates = [t for t in metrics["templates"] if t["year"] == cur]          # one row per send
    dist = M.distinct(metrics, cur)                                                # one entry per template
    for k, (name, rows_) in enumerate(dist.items(), 1):
        t = rows_[0]
        st = M.card_stats(rows_)
        a, rq = ana_t[name], req_t[name]
        card = rq.get("card", {})
        desc = card.get("description", "")
        desc = " ".join(desc) if isinstance(desc, list) else desc
        D.h2(f"{sec4}.{k} {name}: {a['content_risk'].lower()} content risk" + (f" ({rq['label']})" if rq.get("label") and rq["label"] != name else ""))
        wl, wr = 3900, PAGE_W - 3900
        tb = D.d.add_table(rows=0, cols=2)
        _tbl_setup(tb, [wl, wr])
        hr = tb.add_row()
        _row_flags(hr)
        write_cell(hr.cells[0], f"**{name}**   |   {st['when']}   |   GoogleID {t['google_id_short']}", size=8.5, bold=True, color="FFFFFF", fill=hf, width=wl)
        write_cell(hr.cells[1], f"DeliveryRate **{M.fmt_pct(st['dr'])}**   |   ReadRate **{M.fmt_pct(st['rr'])}**   |   Click rate **{M.fmt_pct(st['cr'])}**   |   Content risk: **{a['content_risk']}**",
                   size=8.5, fill=_risk_fill(a["content_risk"]), width=wr)
        for c_ in hr.cells:
            for p_ in c_.paragraphs:
                p_.paragraph_format.keep_with_next = True
        br = tb.add_row()
        _row_flags(br)
        left = []
        img = rq.get("image")
        if img and os.path.exists(os.path.join(image_dir, img)):
            left.append({"picture": os.path.join(image_dir, img), "width": 2.4})
        left += ["**Card text**", f"Title: \"{iconify(card.get('title', ''))}\" Description: \"{iconify(desc)}\" Button: {iconify(card.get('button', ''))}"]
        right = [f"**Creative:** {rq.get('creative', '')}", "**Works**"] + [{"bullet": f"**{x['h']}.** {x['b']}"} for x in a.get("works", [])] + \
                ["**Watch-outs**"] + [{"bullet": f"**{x['h']}.** {x['b']}"} for x in a.get("watch_outs", [])]
        write_cell(br.cells[0], left, size=7.5, width=wl)
        write_cell(br.cells[1], right, size=8.5, width=wr)
        D.gap()

    k = len(dist) + 1
    D.h2(f"{sec4}.{k} Click rate by send ({len(dist)} templates, {len(cur_templates)} sends, {ylabel(cur)})")
    D.p("P2AResponseRate is the sheet's click rate (P2AResponses \u00F7 Messages Read). The creative and text are only seen by people who read the message, so the read base is the right measure for judging them. The delivered base is shown for an end-to-end view (marked *, our calculation).")
    g = metrics["groups"]["year"][str(cur)]
    rows = []
    for t in sorted(cur_templates, key=lambda t: -t["rates"]["P2AResponseRate"]):
        idx = t["click_index_vs_year"] or 0
        rows.append([t["template"], M.nice_date(t["date"]).rsplit(" ", 1)[0], M.fmt_int(t["counts"]["read"]), M.fmt_int(t["counts"]["p2a_responses"]),
                     {"t": "**" + M.fmt_pct(M.shown_rate(t, "P2AResponseRate")) + "**", "fill": "green" if idx >= 1 else ("amber" if idx >= 0.75 else "red")},
                     f"{idx:.2f}x", M.fmt_pct(t["rates"]["click_rate_delivered_base"])])
    rows.append([{"t": f"**Total ({len(cur_templates)} sends)**", "fill": "grey"}, {"t": "", "fill": "grey"}, {"t": f"**{M.fmt_int(g['counts']['read'])}**", "fill": "grey"},
                 {"t": f"**{M.fmt_int(g['counts']['p2a_responses'])}**", "fill": "grey"}, {"t": f"**{M.fmt_pct(g['rates']['P2AResponseRate'])}**", "fill": "grey"},
                 {"t": "**1.00x**", "fill": "grey"}, {"t": f"**{M.fmt_pct(g['rates']['click_rate_delivered_base'])}**", "fill": "grey"}])
    add_table(D.d, [2550, 800, 1150, 1150, 1350, 1346, 1400], ["Template", "Date", "Messages Read", "P2AResponses", "P2AResponseRate (click rate)", "Index vs blended click rate*", "Click rate, delivered base*"], rows, size=8, header_fill=hf)
    D.note(f"Index = click rate of the send \u00F7 blended click rate of all {len(cur_templates)} sends ({M.fmt_pct(g['rates']['P2AResponseRate'])}).")

    fm = analysis.get("feature_matrix")
    if fm and fm.get("rows"):
        D.p("**Which text features line up with click rate**", keep=True, after=3)
        nc = len(fm["header"])
        mid = (PAGE_W - 1500 - 2046) // max(1, nc - 2)
        w = [1500] + [mid] * (nc - 2) + [PAGE_W - 1500 - mid * (nc - 2)]
        add_table(D.d, w, fm["header"], fm["rows"], size=7.5, header_fill=hf)
        D.gap()
        for b in fm.get("bullets", []):
            D.bullet(b)

    refs = [r for r in (request.get("reference_templates") or []) if r["template"] in tpl]
    if refs:
        k += 1
        D.h2(f"{sec4}.{k} Earlier-period templates")
        D.p("No graphics are available for these templates. The text below is reproduced from the portal.", keep=True)
        rows = []
        for r in refs:
            t = tpl[r["template"]]
            c = r.get("card", {})
            desc = c.get("description", "")
            desc = " ".join(desc) if isinstance(desc, list) else desc
            rows.append([r["template"], M.nice_date(t["date"]), [f"**Title:** \"{iconify(c.get('title', ''))}\"", f"**Description:** \"{iconify(desc)}\"", f"**Button:** {iconify(c.get('button', ''))}"],
                         f"Messages Submitted {M.fmt_int(t['counts']['submitted'])}; DeliveryRate {M.fmt_pct(M.shown_rate(t, 'DeliveryRate'))}; ReadRate {M.fmt_pct(M.shown_rate(t, 'ReadRate'))}; click rate **{M.fmt_pct(M.shown_rate(t, 'P2AResponseRate'))}**"])
        add_table(D.d, [1650, 900, 4200, 2996], ["Template", "Date", "Card text", "Metrics"], rows, size=8, header_fill=hf)
        D.gap()

    for es in analysis.get("extra_sections", []) or []:
        k += 1
        D.h2(f"{sec4}.{k} {es['heading']}")
        if es.get("intro"):
            D.p(es["intro"], keep=True)
        if es.get("rows"):
            add_table(D.d, ratio_widths(*([1.4] + [1] * (len(es["header"]) - 1))), es["header"], es["rows"], size=8, header_fill=hf)
            D.gap()
        for b in es.get("bullets", []):
            D.bullet(b)

    D.h1(f"{nxt()}. Cross-template findings")
    D.numbered(_hb(analysis["findings"]))

    D.h1(f"{nxt()}. Root-cause hypotheses")
    D.p("Where each candidate explanation stands on the evidence so far.", keep=True)
    rows = [[f"{h['id']}. {h['hypothesis']}", h["evidence"], {"t": h["status"], "fill": _status_fill(h["status_class"])}, h["next_step"]] for h in analysis["hypotheses"]]
    add_table(D.d, [2100, 4100, 1600, 1946], ["Hypothesis", "Evidence", "Status", "Next step and owner"], rows, size=8, header_fill=hf)

    if analysis.get("questions"):
        D.h1(f"{nxt()}. {brand}'s questions")
        add_table(D.d, [3000, 4300, 2446], ["Question", "Our view", "Needed"], [[q["question"], q["view"], q["needed"]] for q in analysis["questions"]], size=8.5, header_fill=hf)

    D.h1(f"{nxt()}. Recommendations")
    rows = [[str(r["n"]), r["text"], r["why"], r["owner"], {"t": r["priority"], "fill": _prio_fill(r["priority"])}] for r in analysis["recommendations"]]
    add_table(D.d, [450, 3300, 3150, 1500, 1346], ["#", "Recommendation", "Why (evidence)", "Owner", "Priority"], rows, size=8.5, header_fill=hf)
    D.gap()
    if analysis.get("dos") or analysis.get("donts"):
        D.h2(f"Template dos and don'ts (draft for {brand})")
        add_table(D.d, [4873, 4873], ["Do", "Don't"], [[[{"bullet": x} for x in analysis.get("dos", [])], [{"bullet": x} for x in analysis.get("donts", [])]]], size=8.5, header_fill=hf)

    D.h1(f"{nxt()}. Data still needed")
    rows = [[d["owner"], d["item"], d["why"], {"t": d["priority"], "fill": {"High": "red", "Medium": "amber"}.get(d["priority"], "grey")}] for d in analysis["data_needed"]]
    add_table(D.d, [1300, 4300, 3000, 1146], ["Owner", "Item", "Why it matters", "Priority"], rows, size=8.5, header_fill=hf)

    D.save(out_path)
    return out_path
