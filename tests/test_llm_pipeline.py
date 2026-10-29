import io
import json
import os
import sys
import unittest
import zipfile

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
os.environ.setdefault("SOFFICE_CMD", "soffice")
import helpers as H  # noqa: E402
from mock_llm import MockLLM  # noqa: E402
from rcsapp import ASSETS  # noqa: E402
from rcsapp import pipeline as P  # noqa: E402
from rcsapp import training as T  # noqa: E402
from rcsapp.engines import llm_engine as L  # noqa: E402
from rcsapp.knowledge import load_knowledge  # noqa: E402

BUNDLED = os.path.join(ASSETS, "vi_template.pptx")


def make_input(**kw):
    base = dict(sheet_bytes=H.sheet_bytes(), sheet_name="demo.xlsx", brand="DEMO BRAND", brand_display="Demo Brand", templates=[dict(t) for t in H.TEMPLATES],
                period="August 2026 vs August 2025", period_labels={"2026": "August 2026", "2025": "August 2025"}, issue="The brand sees lower click rates on some templates.",
                images=H.images(), engine="ml", vertical="E-commerce", outputs=("docx", "pptx"))
    base.update(kw)
    return P.JobInput(**base)


def slides_text(pptx_bytes):
    from pptx import Presentation
    prs = Presentation(io.BytesIO(pptx_bytes))
    return [" ".join(sh.text_frame.text for sh in s.shapes if sh.has_text_frame) for s in prs.slides]


class Setup(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.ml_res = P.run_job(make_input(), log=lambda *a: None)
        assert cls.ml_res.status == "built", cls.ml_res.issues

    def cfg(self, url, vision=False, key="k"):
        return L.LLMConfig(base_url=url, api_key=key, model="mock-model", vision=vision, retries=0, timeout=20)


class InternalMLPipeline(Setup):
    def test_builds_word_and_deck(self):
        self.assertEqual(sorted(self.ml_res.files), ["Demo_Brand_RCS_Template_Analysis.docx", "Demo_Brand_RCS_Template_Deck.pptx"])

    def test_the_bundled_vi_template_is_used_by_default(self):
        self.assertTrue(os.path.exists(BUNDLED))
        notes = " ".join(i["message"] for i in self.ml_res.issues)
        self.assertIn("bundled Vi template", notes)
        self.assertNotIn("plain_deck", [i["code"] for i in self.ml_res.issues])
        from pptx import Presentation
        prs = Presentation(io.BytesIO(self.ml_res.files["Demo_Brand_RCS_Template_Deck.pptx"]))
        self.assertIn("Only Title slide", {s.slide_layout.name.strip() for s in prs.slides})

    def test_an_uploaded_template_overrides_the_bundled_one(self):
        r = P.run_job(make_input(template_bytes=open(BUNDLED, "rb").read(), outputs=("pptx",)), log=lambda *a: None)
        self.assertEqual(r.status, "built")
        self.assertIn("the template you uploaded", " ".join(i["message"] for i in r.issues))

    def test_a_pptx_that_is_not_the_vi_template_falls_back_with_a_warning(self):
        from pptx import Presentation
        bio = io.BytesIO()
        Presentation().save(bio)
        r = P.run_job(make_input(template_bytes=bio.getvalue(), outputs=("pptx",)), log=lambda *a: None)
        self.assertEqual(r.status, "built")
        self.assertIn("template_rejected", [i["code"] for i in r.issues])

    def test_pdf_is_made(self):
        r = P.run_job(make_input(outputs=("pdf",)), log=lambda *a: None)
        self.assertTrue(next(iter(r.files.values())).startswith(b"%PDF"))

    def test_emoji_show_as_icon_placeholders_in_the_files(self):
        tpls = [dict(t) for t in H.TEMPLATES]
        tpls[0]["title"] = "Flat \u20B9100 off \U0001F389 with code SAVE100"
        r = P.run_job(make_input(templates=tpls, outputs=("docx", "pptx")), log=lambda *a: None)
        self.assertEqual(r.status, "built")
        from docx import Document
        d = Document(io.BytesIO(r.files["Demo_Brand_RCS_Template_Analysis.docx"]))
        text = "\n".join(c.text for t in d.tables for row in t.rows for c in row.cells)
        self.assertIn("[icon]", text)
        self.assertNotIn("\U0001F389", text)
        self.assertNotIn("\U0001F389", " ".join(slides_text(r.files["Demo_Brand_RCS_Template_Deck.pptx"])))

    def test_repeated_sends_build_one_card_per_template(self):
        r = P.run_job(make_input(sheet_bytes=H.sheet_bytes(repeat_send=True), outputs=("docx",)), log=lambda *a: None)
        self.assertEqual(r.status, "built", [i["message"] for i in r.issues if i["level"] == "error"])
        from docx import Document
        d = Document(io.BytesIO(next(iter(r.files.values()))))
        self.assertEqual(len([p for p in d.paragraphs if "content risk" in p.text]), 4)

    def test_an_unreadable_image_does_not_stop_the_run(self):
        imgs = H.images()
        imgs["demo_generic.png"] = b"this is not an image"
        r = P.run_job(make_input(images=imgs), log=lambda *a: None)
        self.assertEqual(r.status, "built")
        self.assertIn("image_unreadable", [i["code"] for i in r.issues])

    def test_no_internal_marking_in_the_slides_or_report_text(self):
        for name, data in self.ml_res.files.items():
            with zipfile.ZipFile(io.BytesIO(data)) as z:
                parts = [n for n in z.namelist() if n.startswith(("word/document", "ppt/slides/slide", "ppt/notesSlides"))]
                blob = " ".join(z.read(n).decode("utf8", "ignore") for n in parts)
            for bad in ("Vodafone Idea Internal", "C2 \u2013", "Classification", "rule-based"):
                self.assertNotIn(bad, blob, f"{bad} in {name}")

    def test_numbers_in_word_come_from_the_sheet(self):
        from docx import Document
        d = Document(io.BytesIO(self.ml_res.files["Demo_Brand_RCS_Template_Analysis.docx"]))
        text = "\n".join(p.text for p in d.paragraphs) + "\n".join(c.text for t in d.tables for r in t.rows for c in r.cells)
        self.assertIn("1,000,000", text)
        self.assertIn("5.20%", text)

    def test_preflight_gives_plain_words(self):
        r = P.run_job(make_input(issue="", templates=[{"template": "demo_code_100", "title": "", "description": ""}]), log=lambda *a: None)
        self.assertEqual(r.status, "stopped_inputs")
        msg = " ".join(i["message"] for i in r.issues)
        self.assertIn("no card title or description", msg)
        self.assertIn("issue", msg)

    def test_editing_runs_the_checks_again(self):
        bad = json.loads(json.dumps(self.ml_res.analysis))
        bad["summary"]["lead"] += " This is 99.99% certain."
        r = P.run_job(make_input(), log=lambda *a: None, analysis_override=bad)
        self.assertEqual((r.status, r.files), ("blocked_by_checks", {}))
        good = json.loads(json.dumps(self.ml_res.analysis))
        good["summary"]["lead"] += " A reviewer added this sentence."
        self.assertEqual(P.run_job(make_input(), log=lambda *a: None, analysis_override=good).status, "built")

    def test_the_bundled_history_travels_with_the_run_and_comes_back_extended(self):
        h, _ = T.load_seed()
        r = P.run_job(make_input(history=h), log=lambda *a: None)
        self.assertEqual(r.status, "built")
        back = T.read_history(r.history_csv)
        self.assertEqual(len(back), len(h) + 6)
        self.assertIn("Demo Brand", set(back["brand"]))

    def test_every_brand_in_the_history_is_blocked_automatically(self):
        h, _ = T.load_seed()
        good = json.loads(json.dumps(self.ml_res.analysis))
        good["findings"][0]["b"] += " Myntra saw the same."
        r = P.run_job(make_input(history=h, extra_blocked_brands=T.brand_names(h)), log=lambda *a: None, analysis_override=good)
        self.assertEqual(r.status, "blocked_by_checks")
        self.assertIn("Myntra", " ".join(i["message"] for i in r.issues))

    def test_manual_mode_no_longer_exists(self):
        self.assertFalse(hasattr(P, "prepare_brief"))
        self.assertFalse(hasattr(L, "parse_manual"))
        self.assertEqual(set(P.ENGINE_LABELS), {"ml", "llm", "both"})


class HostedLLM(Setup):
    def run_llm(self, mode="good", engine="llm", consent=True, vision=False, key="k"):
        with MockLLM(self.ml_res.analysis, mode) as srv:
            r = P.run_job(make_input(engine=engine, llm_cfg=self.cfg(srv.url, vision, key), llm_consent=consent, extra_blocked_brands=["Axis"]), log=lambda *a: None)
            return r, srv

    def test_two_step_flow_files_and_the_key_is_sent(self):
        r, srv = self.run_llm(key="secret-key")
        self.assertEqual(r.status, "built", [i["message"] for i in r.issues if i["level"] == "error"])
        self.assertEqual(len(srv.calls), 5)
        self.assertTrue(all("STEP 1" in c["first"] for c in srv.calls[:4]))
        self.assertIn("STEP 2", srv.calls[4]["first"])
        self.assertTrue(all(c["auth"] == "Bearer secret-key" for c in srv.calls))

    def test_step1_prompt_has_no_performance_numbers(self):
        from rcsapp import brief as B
        req = {"brand": "DEMO BRAND", "brand_display": "Demo", "period": "", "issue": "x", "vertical": ""}
        t = {"template": "demo_code_100", "creative": "banner", "card": {"title": "Flat off", "description": "d", "button": "b"}}
        p = B.step1_prompt(req, t, load_knowledge())
        for banned in ("PERFORMANCE", "| Template |", "5.20%", "1,000,000", "Totals by year"):
            self.assertNotIn(banned, p)

    def test_wrong_number_is_sent_back_and_corrected(self):
        r, srv = self.run_llm("bad_number_once")
        self.assertEqual(r.status, "built")
        self.assertEqual([c["repair"] for c in srv.calls[4:]], [False, True])

    def test_fenced_answers_and_reasoning_tags_are_parsed(self):
        self.assertEqual(self.run_llm("fenced")[0].status, "built")
        self.assertEqual(self.run_llm("think")[0].status, "built")
        self.assertEqual(L.extract_json('<think>maybe {"a": 0}</think> {"a": 1}'), {"a": 1})

    def test_other_brand_in_answer_blocks_the_run(self):
        r, _ = self.run_llm("other_brand")
        self.assertNotEqual(r.status, "built")
        self.assertEqual(r.files, {})
        self.assertIn("Axis", " ".join(i["message"] for i in r.issues))

    def test_endpoint_down_and_bad_key_are_plain_errors_not_crashes(self):
        r, _ = self.run_llm("down")
        self.assertIn("model endpoint", " ".join(i["message"] for i in r.issues))
        self.assertEqual(r.files, {})
        r, _ = self.run_llm("bad_key")
        self.assertIn("401", " ".join(i["message"] for i in r.issues))

    def test_a_rate_limit_is_retried(self):
        with MockLLM(self.ml_res.analysis, "rate_limit_once") as srv:
            cfg = L.LLMConfig(base_url=srv.url, api_key="k", model="m", retries=2, timeout=20)
            self.assertTrue(L.ping(cfg)[0])
            self.assertEqual(len(srv.calls), 2)

    def test_a_provider_that_wants_max_completion_tokens_still_works(self):
        with MockLLM(self.ml_res.analysis, "new_token_name") as srv:
            cfg = L.LLMConfig(base_url=srv.url, api_key="k", model="m", retries=1, timeout=20)
            ok, msg = L.ping(cfg)
            self.assertTrue(ok, msg)
            self.assertIn("max_completion_tokens", srv.bodies[-1])
            self.assertNotIn("max_tokens", srv.bodies[-1])

    def test_consent_is_required(self):
        r, srv = self.run_llm(consent=False)
        self.assertEqual((r.status, srv.calls), ("stopped_inputs", []))

    def test_both_mode_flags_disagreement_and_uses_the_cautious_rating(self):
        r, _ = self.run_llm("flip", engine="both")
        self.assertEqual(r.status, "built")
        self.assertFalse(next(x for x in r.agreement if x["template"] == "demo_code_100")["agree"])
        self.assertEqual(next(t for t in r.analysis["templates"] if t["template"] == "demo_code_100")["content_risk"], "High")
        self.assertIn("Internal ML model", r.review_notes)

    def test_both_mode_falls_back_to_the_internal_model_if_the_llm_is_down(self):
        r, _ = self.run_llm("down", engine="both")
        self.assertEqual(r.status, "built")
        self.assertIn("fallback", [i["code"] for i in r.issues])

    def test_images_are_sent_only_when_vision_is_on(self):
        _, off = self.run_llm(vision=False)
        _, on = self.run_llm(vision=True)
        self.assertFalse(any(c["has_image"] for c in off.calls))
        self.assertTrue(any(c["has_image"] for c in on.calls[:4]))

    def test_a_refused_image_falls_back_to_text_and_the_run_still_builds(self):
        r, srv = self.run_llm("refuse_image", vision=True)
        self.assertEqual(r.status, "built", [i["message"] for i in r.issues if i["level"] == "error"])
        self.assertEqual(sum(c["has_image"] for c in srv.calls), 4)
        self.assertEqual(len(srv.calls), 9)
        notes = [i for i in r.issues if i["code"] == "image_not_used"]
        self.assertEqual(len(notes), 4)
        self.assertIn("text alone", notes[0]["message"])

    def test_connection_test_checks_text_and_image_input(self):
        with MockLLM(self.ml_res.analysis, "good") as srv:
            ok, msg = L.ping(self.cfg(srv.url, vision=True), with_image=True)
            self.assertTrue(ok, msg)
            self.assertIn("Image input works", msg)
            self.assertTrue(srv.calls[-1]["has_image"])
        with MockLLM(self.ml_res.analysis, "refuse_image") as srv:
            ok, msg = L.ping(self.cfg(srv.url, vision=True), with_image=True)
            self.assertFalse(ok)
            self.assertIn("image input failed", msg)
        self.assertFalse(L.ping(L.LLMConfig(base_url="", model=""))[0])

    def test_gold_rating_by_the_language_model(self):
        rows = [{"template": "demo_code_100", "brand": "B", "vertical": "E-commerce", "title": "Flat 100 off", "description": "Use code SAVE100"},
                {"template": "demo_raffle", "brand": "B", "vertical": "E-commerce", "title": "Win big", "description": "Stand a chance to win"}]
        with MockLLM(self.ml_res.analysis, "good") as srv:
            out = L.rate_content(self.cfg(srv.url), rows, load_knowledge(), {"blocked_phrases": []})
        self.assertIn(out["demo_code_100"]["content_risk"], ("Low", "Medium", "High"))
        self.assertEqual(len(out), 2)

    def test_config_from_env_and_secrets(self):
        c = L.config_from({"base_url": "https://api.groq.com/openai/v1", "model": "qwen/qwen3.8-27b", "vision": "true"}, env={})
        self.assertTrue(c.ready() and c.vision)
        self.assertEqual(c.host, "api.groq.com")
        self.assertFalse(L.config_from({}, env={}).ready())
        self.assertEqual(L.config_from({}, env={"LLM_BASE_URL": "http://localhost:11434/v1", "LLM_MODEL": "llama3.2-vision"}).host, "localhost:11434")


if __name__ == "__main__":
    unittest.main()
