"""Synthetic data for the tests. Nothing here is real brand data."""
import io

import numpy as np
import openpyxl
import pandas as pd
from PIL import Image

HDR = ["Date", "Aggregator", "SubAggregator", "Brand", "AgentID", "GoogleID", "MessageType", "AgentName", "Template", "Messages Submitted", "Messages Sent",
       "Messages Delivered", "Messages Read", "Messages Failed", "Messages Expired", "Messages DeliveryRate", "ReadRate", "P2AResponses=clicks",
       "P2AResponseRate=clickrate", "P2AMessages", "template title", "template body", "CTA", "CTA URL", "vertical", "exact text"]
FOOTER = "&8C1 \u2013 Vodafone Idea External"


def _row(date, brand, gid, tpl, sub, sent, dl, rd, p2a, title, body, cta, exact=""):
    fail, exp = sub - sent, sent - dl
    return [date, "PartnerX", "-", brand, "id1", gid, "Promo", "Demo", tpl, sub, sent, dl, rd, fail, exp,
            f"{dl / sent * 100:.2f} %", f"{rd / dl * 100:.2f} %", p2a, f"{p2a / rd * 100:.2f} %", 0, title, body, cta, "http://x.in/[0_sh_code]", "apparel", exact]


TEXTS = {
    "demo_code_100": ("Flat \u20B9100 off with code SAVE100", "Use code SAVE100 on orders above \u20B9499. Ends tonight.", "Shop Now"),
    "demo_generic": ("Big sale is here", "Shop now for great deals on your favourites.", "Shop Now"),
    "demo_raffle": ("[v1], you're eligible for the [v2] Giveaway!", "It's your LUCKY DAY! Stand a chance to win exclusive gifts. WIN BIG! Hurry, the clock's ticking! Shop for \u20B93,000+ between 4 and 6 PM to be considered for the draw this evening only.", "Shop Now"),
    "demo_flat_min": ("PRICE DROP ALERT: MIN 30% OFF!", "[v1], grab MIN 30% OFF on top picks! Hurry!", "Shop Now"),
    "old_code_50": ("[v2], you've unlocked a coupon", "Get [v3] OFF with code [v4]. Valid only till [v5].", "Shop Now"),
    "old_generic": ("Festive picks for you", "Great gifts for everyone you love.", "Shop now"),
}


def sheet_bytes(brand="DEMO BRAND LTD", extra_brand=False, plain_headers=False, footer=True, with_text=True, repeat_send=False):
    wb = openpyxl.Workbook()
    ws = wb.active
    hdr = list(HDR)
    if plain_headers:
        hdr[17], hdr[18] = "P2AResponses", "P2AResponseRate"
    ws.append(hdr)
    g = "demo_abc123_agent"
    rows = [
        ("2026-08-05", "demo_code_100", 1000000, 600000, 540000, 100000, 5200),
        ("2026-08-12", "demo_generic", 1200000, 700000, 630000, 118000, 1900),
        ("2026-08-19", "demo_raffle", 900000, 500000, 440000, 90000, 800),
        ("2026-08-26", "demo_flat_min", 800000, 480000, 430000, 85000, 1500),
        ("2025-08-06", "old_code_50", 700000, 420000, 410000, 60000, 4200),
        ("2025-08-14", "old_generic", 600000, 360000, 350000, 52000, 1100),
    ]
    if repeat_send:
        rows.append(("2026-08-30", "demo_code_100", 900000, 540000, 490000, 90000, 3600))
    for d, t, sub, sent, dl, rd, p2a in rows:
        ti, bo, cta = TEXTS[t] if with_text else ("", "", "")
        ws.append(_row(d, brand, g, t, sub, sent, dl, rd, p2a, ti, (ti + " " + bo).strip() if with_text else "", cta))
    ws.append([None] * 26)
    ws.append([2025] + [None] * 25)
    if extra_brand:
        ws.append(_row("2026-08-09", "OTHER BRAND LTD", "oth_x_agent", "oth_t1", 500000, 300000, 270000, 50000, 900, "Other", "Other text", "Go"))
    if footer:
        last = ws.max_row
        ws.cell(row=last, column=1).value = FOOTER if ws.cell(row=last, column=1).value is None else str(ws.cell(row=last, column=1).value) + FOOTER
    bio = io.BytesIO()
    wb.save(bio)
    return bio.getvalue()


TEMPLATES = [
    {"template": t, "label": t, "short": t, "creative": "A banner.", "title": TEXTS[t][0], "description": TEXTS[t][1], "button": TEXTS[t][2]}
    for t in ["demo_code_100", "demo_generic", "demo_raffle", "demo_flat_min", "old_code_50", "old_generic"]
]


def png_bytes(color=(200, 60, 60), size=(800, 400)):
    bio = io.BytesIO()
    Image.new("RGB", size, color).save(bio, format="PNG")
    return bio.getvalue()


def images():
    return {t["template"] + ".png": png_bytes((40 + 30 * i, 80, 160)) for i, t in enumerate(TEMPLATES[:4])}


WORDS = ["river", "stone", "cloud", "amber", "violet", "marble", "harbor", "meadow", "copper", "lantern", "velvet", "thunder", "silver", "orchid", "canyon",
         "falcon", "breeze", "crystal", "ember", "garden", "island", "jungle", "kitten", "lemon", "mountain", "nectar", "ocean", "pepper", "quartz", "rocket",
         "saddle", "tundra", "umbrella", "valley", "willow", "yonder", "zephyr", "anchor", "bridge", "candle"]


def synthetic_history_raw(n_brands=6, per=14, signal=True, seed=0, cohort=""):
    """Raw history rows (before T.finalize) with a planted effect of the text on the click rate."""
    rng = np.random.default_rng(seed)
    rows = []
    for b in range(n_brands):
        for i in range(per):
            num = int(rng.integers(0, 2))
            chance = int(rng.random() < 0.4)
            code = int(rng.integers(0, 2))
            title = f"Get \u20B9{int(rng.integers(50, 500))} off" if num else "Great deals here"
            desc = " ".join(rng.choice(WORDS, 9, replace=False)) + (" Stand a chance to win" if chance else "") + (f" Use code SAVE{i}" if code else "")
            effect = (0.7 * num - 0.8 * chance + 0.3 * code + rng.normal(0, 0.12)) if signal else rng.normal(0, 0.6)
            reads = 50000
            clicks = int(round(reads * 0.012 * np.exp(effect)))
            rows.append({"brand": f"Brand{b}", "vertical": "E-commerce", "cohort": cohort, "template": f"b{b}_t{i}", "date": f"2026-{1 + i % 9:02d}-{1 + (i * 3) % 27:02d}",
                         "reads": reads, "clicks": clicks, "title": title, "description": desc, "button": "Shop Now", "exclude": "", "exclude_reason": "", "source": "synthetic"})
    return pd.DataFrame(rows)
